import hashlib
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
import xlrd
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.core.exceptions import SmartOpsError
from app.models.conversation import utc_now
from app.models.files import AllowedPath, DocumentChunk, FileAuditLog, FileWriteRequest, IndexedFile
from app.schemas.files import AllowedPathCreate, FileSearchResult, WriteRequestCreate
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from app.rag.service import RAGService

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log", ".xml", ".yaml", ".yml"}
CONTENT_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".pptx"}


class FileAccessError(SmartOpsError):
    """Raised when a file operation violates an allow-list or approval rule."""


class FileAccessService:
    def __init__(self, db: Session, settings: Settings, rag: "RAGService | None" = None) -> None:
        self.db = db
        self.settings = settings
        self.rag = rag

    def list_allowed_paths(self) -> list[AllowedPath]:
        return list(self.db.scalars(select(AllowedPath).order_by(AllowedPath.created_at.desc())))

    def add_allowed_path(self, payload: AllowedPathCreate) -> AllowedPath:
        root = Path(payload.root_path).expanduser().resolve(strict=True)
        if not root.is_dir():
            raise FileAccessError("İzin verilen yol mevcut bir klasör olmalıdır")
        existing = self.db.scalar(select(AllowedPath).where(AllowedPath.root_path == str(root)))
        if existing:
            return existing
        item = AllowedPath(label=payload.label, root_path=str(root), permission=payload.permission.value)
        self.db.add(item)
        self._audit("allowed_path_added", "admin", str(root), payload.permission.value)
        self.db.commit()
        self.db.refresh(item)
        return item

    def index_path(self, allowed_path_id: str) -> tuple[int, int, list[str]]:
        allowed = self._get_allowed(allowed_path_id)
        root = Path(allowed.root_path).resolve(strict=True)
        seen: set[str] = set()
        indexed = 0
        skipped = 0
        errors: list[str] = []
        for current_root, directories, files in os.walk(root, followlinks=False):
            directories[:] = [name for name in directories if not name.startswith(".smartops-")]
            for name in files:
                path = Path(current_root, name)
                try:
                    resolved = path.resolve(strict=True)
                    self._assert_within(root, resolved)
                    if resolved.stat().st_size > self.settings.max_index_file_bytes:
                        skipped += 1
                        continue
                    content = self._extract_text(resolved) if resolved.suffix.lower() in CONTENT_EXTENSIONS else ""
                    stat = resolved.stat()
                    absolute = str(resolved)
                    seen.add(absolute)
                    record = self.db.scalar(select(IndexedFile).where(IndexedFile.absolute_path == absolute))
                    values = {
                        "allowed_path_id": allowed.id,
                        "relative_path": str(resolved.relative_to(root)),
                        "file_name": resolved.name,
                        "extension": resolved.suffix.lower(),
                        "size_bytes": stat.st_size,
                        "modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc),
                        "content": content[:250_000],
                        "indexed_at": utc_now(),
                    }
                    if record:
                        for key, value in values.items():
                            setattr(record, key, value)
                    else:
                        record = IndexedFile(absolute_path=absolute, **values)
                        self.db.add(record)
                    self.db.flush()
                    if self.rag and content:
                        self.rag.index_file(record)
                    indexed += 1
                except Exception as exc:  # individual files must not abort a full scan
                    skipped += 1
                    if len(errors) < 20:
                        errors.append(f"{name}: {type(exc).__name__}")
        stale = list(self.db.scalars(select(IndexedFile).where(IndexedFile.allowed_path_id == allowed.id)))
        for record in stale:
            if record.absolute_path not in seen:
                self.db.delete(record)
        self._audit("path_indexed", "admin", str(root), f"indexed={indexed}, skipped={skipped}")
        self.db.commit()
        return indexed, skipped, errors

    def disable_allowed_path(self, allowed_path_id: str) -> AllowedPath:
        allowed = self._get_allowed(allowed_path_id)
        allowed.enabled = False
        self.db.execute(delete(DocumentChunk).where(DocumentChunk.allowed_path_id == allowed.id))
        self._audit("allowed_path_disabled", "admin", allowed.root_path, allowed.permission)
        self.db.commit()
        self.db.refresh(allowed)
        return allowed

    def search(self, query: str, allowed_path_id: str | None = None) -> list[FileSearchResult]:
        terms = [term.casefold() for term in query.split() if len(term) >= 2]
        if not terms:
            return []
        statement = select(IndexedFile).join(AllowedPath).where(AllowedPath.enabled.is_(True))
        if allowed_path_id:
            self._get_allowed(allowed_path_id)
            statement = statement.where(IndexedFile.allowed_path_id == allowed_path_id)
        records = list(self.db.scalars(statement))
        ranked: list[tuple[int, IndexedFile]] = []
        for record in records:
            allowed = self.db.get(AllowedPath, record.allowed_path_id)
            if not allowed:
                continue
            try:
                root = Path(allowed.root_path).resolve(strict=True)
                source = Path(record.absolute_path).resolve(strict=True)
                self._assert_within(root, source)
                if not source.is_file():
                    continue
            except (OSError, FileAccessError):
                continue
            haystack = f"{record.file_name}\n{record.relative_path}\n{record.content}".casefold()
            score = sum(4 if term in record.file_name.casefold() else 1 for term in terms if term in haystack)
            if score:
                ranked.append((score, record))
        ranked.sort(key=lambda pair: (pair[0], pair[1].modified_at), reverse=True)
        return [self._search_result(record, terms) for _, record in ranked[: self.settings.max_search_results]]

    def search_for_assistant(self, query: str) -> dict:
        matches = self.search(query)
        return {
            "query": query,
            "matches": [
                {
                    "indexed_file_id": match.indexed_file_id,
                    "file_name": match.file_name,
                    "relative_path": match.relative_path,
                    "directory_path": str(Path(match.absolute_path).parent),
                    "allowed_path_id": self._allowed_path_id_for(match.absolute_path),
                    "excerpt": match.excerpt,
                }
                for match in matches
            ],
            "instruction": "Yanıtı yalnızca eşleşen belge parçalarına dayandır; kullanılan dosya adlarını ve klasör yollarını kaynak olarak belirt.",
        }

    def create_write_request(self, payload: WriteRequestCreate) -> FileWriteRequest:
        allowed = self._get_allowed(payload.allowed_path_id)
        if allowed.permission != "read_write":
            raise FileAccessError("Bu kök yalnızca okuma yetkisine sahip")
        target = self._safe_target(allowed, payload.relative_path)
        if not target.parent.exists():
            raise FileAccessError("Hedef üst klasör mevcut değil")
        expected_hash = self._sha256(target) if target.exists() else None
        operation = "replace" if target.exists() else "create"
        request = FileWriteRequest(
            allowed_path_id=allowed.id,
            relative_path=str(target.relative_to(Path(allowed.root_path))),
            operation=operation,
            proposed_content=payload.content,
            expected_sha256=expected_hash,
            requested_by=payload.requested_by,
        )
        self.db.add(request)
        self._audit("write_requested", payload.requested_by, str(target), operation)
        self.db.commit()
        self.db.refresh(request)
        return request

    def resolve_indexed_file(self, indexed_file_id: str) -> Path:
        if not indexed_file_id.isdigit():
            raise FileAccessError("Geçersiz dosya kimliği")
        record = self.db.get(IndexedFile, int(indexed_file_id))
        if not record:
            raise FileAccessError("İndekslenmiş dosya bulunamadı")
        allowed = self._get_allowed(record.allowed_path_id)
        root = Path(allowed.root_path).resolve(strict=True)
        path = Path(record.absolute_path).resolve(strict=True)
        self._assert_within(root, path)
        if not path.is_file():
            raise FileAccessError("Dosya artık mevcut değil")
        return path

    def list_write_requests(self) -> list[FileWriteRequest]:
        return list(self.db.scalars(select(FileWriteRequest).order_by(FileWriteRequest.created_at.desc())))

    def approve_write(self, request_id: UUID, actor: str) -> FileWriteRequest:
        request = self._get_write_request(request_id)
        if request.status != "pending":
            raise FileAccessError("Yalnızca bekleyen talepler onaylanabilir")
        allowed = self._get_allowed(request.allowed_path_id)
        if allowed.permission != "read_write":
            raise FileAccessError("Yazma yetkisi kaldırılmış")
        target = self._safe_target(allowed, request.relative_path)
        current_hash = self._sha256(target) if target.exists() else None
        if current_hash != request.expected_sha256:
            raise FileAccessError("Dosya talep oluşturulduktan sonra değişmiş; güvenlik için işlem durduruldu")
        if target.exists():
            backup_dir = Path(allowed.root_path, ".smartops-backups")
            backup_dir.mkdir(exist_ok=True)
            stamp = utc_now().strftime("%Y%m%dT%H%M%SZ")
            shutil.copy2(target, backup_dir / f"{target.name}.{stamp}.bak")
        temporary = target.with_name(f".{target.name}.smartops-tmp")
        temporary.write_text(request.proposed_content, encoding="utf-8")
        os.replace(temporary, target)
        request.status = "approved"
        request.approved_by = actor
        request.decided_at = utc_now()
        self._audit("write_approved", actor, str(target), request.operation)
        self.db.commit()
        self.db.refresh(request)
        return request

    def reject_write(self, request_id: UUID, actor: str) -> FileWriteRequest:
        request = self._get_write_request(request_id)
        if request.status != "pending":
            raise FileAccessError("Yalnızca bekleyen talepler reddedilebilir")
        allowed = self._get_allowed(request.allowed_path_id)
        target = self._safe_target(allowed, request.relative_path)
        request.status = "rejected"
        request.approved_by = actor
        request.decided_at = utc_now()
        self._audit("write_rejected", actor, str(target), request.operation)
        self.db.commit()
        self.db.refresh(request)
        return request

    def audit_logs(self) -> list[FileAuditLog]:
        return list(self.db.scalars(select(FileAuditLog).order_by(FileAuditLog.created_at.desc()).limit(200)))

    def _get_allowed(self, allowed_path_id: str) -> AllowedPath:
        item = self.db.get(AllowedPath, allowed_path_id)
        if not item or not item.enabled:
            raise FileAccessError("İzin verilen klasör bulunamadı veya devre dışı")
        return item

    def _get_write_request(self, request_id: UUID) -> FileWriteRequest:
        request = self.db.get(FileWriteRequest, str(request_id))
        if not request:
            raise FileAccessError("Yazma talebi bulunamadı")
        return request

    def _safe_target(self, allowed: AllowedPath, relative_path: str) -> Path:
        if Path(relative_path).is_absolute():
            raise FileAccessError("Hedef yol izinli köke göre bağıl olmalıdır")
        root = Path(allowed.root_path).resolve(strict=True)
        target = (root / relative_path).resolve(strict=False)
        self._assert_within(root, target)
        if target.suffix.lower() not in TEXT_EXTENSIONS:
            raise FileAccessError("İlk sürüm yalnızca metin tabanlı dosyalara yazabilir")
        return target

    @staticmethod
    def _assert_within(root: Path, target: Path) -> None:
        try:
            common = os.path.commonpath([os.path.normcase(str(root)), os.path.normcase(str(target))])
        except ValueError as exc:
            raise FileAccessError("Hedef izin verilen kökün dışında") from exc
        if common != os.path.normcase(str(root)):
            raise FileAccessError("Hedef izin verilen kökün dışında")

    @staticmethod
    def _extract_text(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in TEXT_EXTENSIONS:
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".pdf":
            return "\n".join(f"Sayfa: {number}\n{page.extract_text() or ''}" for number, page in enumerate(PdfReader(str(path)).pages, start=1))
        if suffix == ".docx":
            return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
        if suffix in {".xlsx", ".xlsm"}:
            workbook = load_workbook(filename=path, read_only=True, data_only=True)
            lines: list[str] = []
            try:
                for sheet in workbook.worksheets:
                    lines.append(f"Sayfa: {sheet.title}")
                    for row in sheet.iter_rows():
                        values = [str(cell.value).strip() for cell in row if cell.value is not None and str(cell.value).strip()]
                        if values:
                            lines.append(" | ".join(values))
                        if sum(len(line) for line in lines) >= 250_000:
                            return "\n".join(lines)[:250_000]
            finally:
                workbook.close()
            return "\n".join(lines)
        if suffix == ".xls":
            workbook = xlrd.open_workbook(path, on_demand=True)
            lines = []
            try:
                for sheet in workbook.sheets():
                    lines.append(f"Sayfa: {sheet.name}")
                    for row_index in range(sheet.nrows):
                        values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
                        if values:
                            lines.append(" | ".join(values))
                        if sum(len(line) for line in lines) >= 250_000:
                            return "\n".join(lines)[:250_000]
            finally:
                workbook.release_resources()
            return "\n".join(lines)
        if suffix == ".pptx":
            presentation = Presentation(str(path))
            lines = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                lines.append(f"Slayt: {slide_number}")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        lines.append(shape.text.strip())
            return "\n".join(lines)
        return ""

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(65_536), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _search_result(record: IndexedFile, terms: list[str]) -> FileSearchResult:
        content = record.content.replace("\n", " ")
        folded = content.casefold()
        positions = [folded.find(term) for term in terms if folded.find(term) >= 0]
        start = max(0, (min(positions) if positions else 0) - 100)
        excerpt = content[start : start + 360].strip()
        return FileSearchResult(
            indexed_file_id=str(record.id),
            file_name=record.file_name,
            relative_path=record.relative_path,
            absolute_path=record.absolute_path,
            extension=record.extension,
            modified_at=record.modified_at,
            excerpt=excerpt,
        )

    def _audit(self, action: str, actor: str, target: str, detail: str) -> None:
        self.db.add(FileAuditLog(action=action, actor=actor, target_path=target, detail=detail))

    def _allowed_path_id_for(self, absolute_path: str) -> str:
        record = self.db.scalar(select(IndexedFile).where(IndexedFile.absolute_path == absolute_path))
        return record.allowed_path_id if record else ""
